import xlsxwriter
from xlsxwriter.workbook import Workbook
import openpyxl
import os
from tqdm import tqdm

class XLSXReader:
    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = openpyxl.load_workbook(file_path, read_only=True)
        
    def get_sheet_data(self, sheet_name=None, chunk_size=1000):
        sheet = self.workbook[sheet_name] if sheet_name else self.workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(row)
            if len(rows) >= chunk_size:
                yield rows
                rows = []
        if rows:
            yield rows

class XLSXWriter:
    def __init__(self, file_path):
        self.workbook = Workbook(file_path, {
            'constant_memory': True,
            'default_date_format': 'yyyy-mm-dd',
            'remove_timezone': True
        })
        
    def write_chunk(self, data, sheet_name='Sheet1'):
        worksheet = self.workbook.add_worksheet(sheet_name)
        for row_num, row_data in enumerate(data):
            worksheet.write_row(row_num, 0, row_data)
        
    def close(self):
        self.workbook.close()


def detect_file_type(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.csv':
        return 'csv'
    elif ext in ('.xls', '.xlsx', '.xlsm', '.xlsb'):
        return 'xlsx'
    raise ValueError(f'不支持的格式: {ext}')